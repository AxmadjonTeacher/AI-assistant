import asyncio
import json
import os
import socket
import subprocess
import threading
import time
import uuid
import ssl
import urllib.request
import urllib.parse
from datetime import datetime
from dataclasses import dataclass, field
from typing import Optional, Dict, Any, List

from config import config

def _update_hud_working(title: str, subtitle: str):
    try:
        from agent_hud import get_agent_hud
        hud = get_agent_hud()
        if hud:
            hud.show_working(title, subtitle)
    except Exception:
        pass

def _update_hud_completed(title: str, subtitle: str, auto_hide_seconds: float = 4.0):
    try:
        from agent_hud import get_agent_hud
        hud = get_agent_hud()
        if hud:
            hud.show_completed(title, subtitle, auto_hide_seconds)
    except Exception:
        pass

def _update_hud_hide():
    try:
        from agent_hud import get_agent_hud
        hud = get_agent_hud()
        if hud:
            hud.hide()
    except Exception:
        pass

@dataclass
class AgentTask:
    task_id: str
    agent_type: str
    title: str
    prompt: str
    status: str = "running" # "running", "completed", "failed"
    started_at: float = field(default_factory=time.time)
    finished_at: Optional[float] = None
    result_message: Optional[str] = None
    error: Optional[str] = None

def resolve_image_file(query: Optional[str] = None, prompt_hint: str = "") -> Optional[str]:
    """Intelligently resolves an image file path from a query or scans the prompt.
    Supports ~/Desktop, ~/Downloads, ~/Pictures, ~/Documents, exact paths, fuzzy names,
    and 'screen' (takes a screenshot).
    """
    import re
    desktop_dir = os.path.expanduser("~/Desktop")
    downloads_dir = os.path.expanduser("~/Downloads")
    pictures_dir = os.path.expanduser("~/Pictures")
    documents_dir = os.path.expanduser("~/Documents")
    search_folders = [desktop_dir, downloads_dir, pictures_dir, documents_dir, os.getcwd()]

    target_str = (query or "").strip()
    if not target_str and prompt_hint:
        # Check quoted filenames
        quoted = re.findall(r'["\']([^"\']+\.(?:png|jpg|jpeg|webp|svg|bmp))["\']', prompt_hint, re.IGNORECASE)
        if quoted:
            target_str = quoted[0]
        else:
            unquoted = re.findall(r'([\w\-_ ]+\.(?:png|jpg|jpeg|webp|svg|bmp))', prompt_hint, re.IGNORECASE)
            if unquoted:
                target_str = unquoted[0].strip()
            elif any(s in prompt_hint.lower() for s in ("screen", "ekran", "screenshot", "active_screen")):
                target_str = "screen"

    if not target_str:
        return None

    target_str = target_str.strip("\"' \t\n")

    # Screen capture check
    if target_str.lower() in ("screen", "screenshot", "active_screen", "ekran", "ekrandagi"):
        try:
            shot_path = f"/tmp/swan_screen_{int(time.time())}.png"
            subprocess.run(["screencapture", "-x", "-C", shot_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            if os.path.isfile(shot_path):
                return shot_path
        except Exception:
            pass

    expanded = os.path.expanduser(target_str)
    if os.path.isfile(expanded):
        return expanded

    clean_name = os.path.basename(target_str)
    clean_norm = clean_name.lower().replace(" ", "").replace("_", "").replace("-", "")

    for folder in search_folders:
        if not os.path.isdir(folder):
            continue
        cand = os.path.join(folder, clean_name)
        if os.path.isfile(cand):
            return cand
        try:
            items = os.listdir(folder)
            for item in items:
                if item.lower() == clean_name.lower():
                    c = os.path.join(folder, item)
                    if os.path.isfile(c): return c
            for item in items:
                item_norm = item.lower().replace(" ", "").replace("_", "").replace("-", "")
                if item_norm == clean_norm:
                    c = os.path.join(folder, item)
                    if os.path.isfile(c): return c
            base_no_ext = os.path.splitext(clean_name)[0].lower()
            if len(base_no_ext) >= 3:
                for item in items:
                    if item.lower().startswith(base_no_ext) and item.lower().endswith(('.png', '.jpg', '.jpeg', '.webp', '.svg')):
                        c = os.path.join(folder, item)
                        if os.path.isfile(c): return c
        except Exception:
            pass

    # Keyword search on Desktop for logo/reference/icon
    if any(kw in target_str.lower() or kw in prompt_hint.lower() for kw in ("logo", "reference", "icon")):
        try:
            dt_items = [f for f in os.listdir(desktop_dir) if f.lower().endswith(('.png', '.jpg', '.jpeg', '.webp', '.svg'))]
            for f in dt_items:
                if "logo" in f.lower() or "reference" in f.lower():
                    return os.path.join(desktop_dir, f)
        except Exception:
            pass

    return None

class AgentManager:
    """Manages background agent tasks and coordinates visual HUD feedback."""
    def __init__(self):
        self.tasks: Dict[str, AgentTask] = {}
        self._lock = threading.Lock()

    def _sync_hud(self, recent_title: Optional[str] = None, is_completion: bool = False):
        """Synchronizes top-right HUD with all concurrent background tasks."""
        with self._lock:
            running = [t for t in self.tasks.values() if t.status == "running"]
            completed = [t for t in self.tasks.values() if t.status == "completed"]

        if is_completion and running:
            # One task completed, but other agents are still actively working!
            # Flash completion for 2.2 seconds without hiding, then revert to the active task(s)
            _update_hud_completed("1 Task Ready ✅", (recent_title or "Task finished")[:25], auto_hide_seconds=0)
            def _resume_working():
                time.sleep(2.2)
                with self._lock:
                    still_running = [t for t in self.tasks.values() if t.status == "running"]
                if still_running:
                    if len(still_running) > 1:
                        names = " + ".join([t.agent_type.replace("_", " ").title() for t in still_running[:2]])
                        _update_hud_working(f"{len(still_running)} Agents Active 🚀", names)
                    else:
                        _update_hud_working("Agent working...", still_running[0].title[:28])
                else:
                    _update_hud_completed("All tasks finished ✅", "", auto_hide_seconds=4.0)
            threading.Thread(target=_resume_working, daemon=True).start()
            return

        if len(running) > 1:
            agent_names = " + ".join([t.agent_type.replace("_", " ").title() for t in running[:2]])
            _update_hud_working(f"{len(running)} Agents Active 🚀", agent_names)
        elif len(running) == 1:
            _update_hud_working("Agent working...", running[0].title[:28])
        else:
            display_title = recent_title or (completed[-1].title if completed else "Tasks completed")
            _update_hud_completed("Agent finished ✅", display_title[:28], auto_hide_seconds=4.0)

    def launch_blender_scene(self, prompt: str, style: str = "cinematic", reference_image: Optional[str] = None) -> Dict[str, Any]:
        """Launches an autonomous 3D director agent to build a scene in Blender,
        or reconstructs a 2D reference logo/image into a high-fidelity 3D model.
        """
        resolved_img = resolve_image_file(reference_image, prompt_hint=prompt)
        is_reconstruction = bool(resolved_img and os.path.isfile(resolved_img))

        short_title = prompt[:30] + ("..." if len(prompt) > 30 else "")
        task_id = f"bld-{uuid.uuid4().hex[:6]}"

        if is_reconstruction:
            img_name = os.path.basename(resolved_img)
            task_title = f"3D Logo: {img_name}"
            task_type = "blender_logo_3d"
        else:
            task_title = f"Blender: {short_title}"
            task_type = "blender"

        task = AgentTask(
            task_id=task_id,
            agent_type=task_type,
            title=task_title,
            prompt=prompt
        )

        with self._lock:
            self.tasks[task_id] = task

        # Update top-right corner HUD with multi-agent awareness
        self._sync_hud()

        # Run execution in background thread so Swan is never blocked
        thread = threading.Thread(target=self._run_blender_worker, args=(task, style, resolved_img), daemon=True)
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": "Blender 3D Vision Artist" if is_reconstruction else "Blender 3D Director",
            "reference_image": resolved_img if is_reconstruction else None,
            "message": f"Blender agent launched to reconstruct '{os.path.basename(resolved_img)}' in 3D." if is_reconstruction else "Blender agent launched in the background. The top-right indicator is active."
        }

    def launch_image_agent(self, prompt: str, source_image_path: Optional[str] = None, aspect_ratio: str = "1:1", model_preference: Optional[str] = None) -> Dict[str, Any]:
        """Launches an autonomous image generation or editing background agent.
        Output is ALWAYS saved directly to the user's Desktop (~/Desktop) and automatically opened.
        """
        is_edit = bool(source_image_path)
        action_name = "Edit Image" if is_edit else "Generate Image"
        short_title = prompt[:30] + ("..." if len(prompt) > 30 else "")
        task_id = f"img-{uuid.uuid4().hex[:6]}"
        task = AgentTask(
            task_id=task_id,
            agent_type="image_edit" if is_edit else "image_gen",
            title=f"{action_name}: {short_title}",
            prompt=prompt
        )

        with self._lock:
            self.tasks[task_id] = task

        # Update top-right corner HUD with multi-agent awareness
        self._sync_hud()

        thread = threading.Thread(
            target=self._run_image_worker,
            args=(task, source_image_path, aspect_ratio, model_preference),
            daemon=True
        )
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": "Swan Visual Artist",
            "output_directory": os.path.expanduser("~/Desktop"),
            "message": f"Image task started in background. The generated image will appear directly on your Desktop (~/Desktop) and open automatically. Visual HUD active."
        }

    def launch_generic_agent(self, agent_type: str, task_description: str, details: str = "") -> Dict[str, Any]:
        """Launches a general background agent for research, scripts, or long tasks."""
        short_title = task_description[:30] + ("..." if len(task_description) > 30 else "")
        task_id = f"agt-{uuid.uuid4().hex[:6]}"
        task = AgentTask(
            task_id=task_id,
            agent_type=agent_type,
            title=f"{agent_type.capitalize()}: {short_title}",
            prompt=task_description
        )

        with self._lock:
            self.tasks[task_id] = task

        self._sync_hud()

        thread = threading.Thread(target=self._run_generic_worker, args=(task, details), daemon=True)
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": agent_type,
            "message": f"{agent_type.capitalize()} agent launched in background. Visual indicator active in top-right corner."
        }

    def launch_transcribe_agent(self, file_path: Optional[str] = None, query: Optional[str] = None) -> Dict[str, Any]:
        """Launches an autonomous audio transcription background agent."""
        task_id = f"aud-{uuid.uuid4().hex[:6]}"
        display_hint = file_path or query or "Latest audio"
        short_title = os.path.basename(display_hint)[:28]
        task = AgentTask(
            task_id=task_id,
            agent_type="transcription",
            title=f"Transcribe: {short_title}",
            prompt=display_hint
        )
        with self._lock:
            self.tasks[task_id] = task

        self._sync_hud()
        thread = threading.Thread(target=self._run_transcribe_worker, args=(task, file_path, query), daemon=True)
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": "Swan Audio Transcriber",
            "message": f"Audio transcription started for '{display_hint}'. Output will be saved to your Desktop and opened automatically."
        }

    def launch_youtube_agent(self, query_or_url: str, focus: str = "") -> Dict[str, Any]:
        """Launches an autonomous YouTube search & executive summarization background agent."""
        task_id = f"yt-{uuid.uuid4().hex[:6]}"
        short_title = query_or_url[:28] + ("..." if len(query_or_url) > 28 else "")
        task = AgentTask(
            task_id=task_id,
            agent_type="youtube_summary",
            title=f"YouTube: {short_title}",
            prompt=query_or_url
        )
        with self._lock:
            self.tasks[task_id] = task

        self._sync_hud()
        thread = threading.Thread(target=self._run_youtube_worker, args=(task, query_or_url, focus), daemon=True)
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": "Swan YouTube Intelligence",
            "message": f"YouTube video summarization started for '{short_title}'. Structured executive summary will be saved to Desktop and opened."
        }

    def launch_document_agent(self, title: str, content: str, format: str = "docx", file_name: str = "") -> Dict[str, Any]:
        """Launches an autonomous document creation background agent (.docx, .pdf, .md)."""
        task_id = f"doc-{uuid.uuid4().hex[:6]}"
        fmt = (format or "docx").lower()
        short_title = title[:28] + ("..." if len(title) > 28 else "")
        task = AgentTask(
            task_id=task_id,
            agent_type=f"doc_{fmt}",
            title=f"Create {fmt.upper()}: {short_title}",
            prompt=content[:100]
        )
        with self._lock:
            self.tasks[task_id] = task

        self._sync_hud()
        thread = threading.Thread(target=self._run_document_worker, args=(task, title, content, fmt, file_name), daemon=True)
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": "Swan Document Publisher",
            "message": f"{fmt.upper()} document generation started for '{title}'. File will appear directly on your Desktop and open automatically."
        }

    def launch_presentation_agent(self, title: str, topic_or_content: str, slide_count: int = 5) -> Dict[str, Any]:
        """Launches an autonomous presentation deck creation background agent (.pptx & web slides)."""
        task_id = f"prs-{uuid.uuid4().hex[:6]}"
        short_title = title[:28] + ("..." if len(title) > 28 else "")
        task = AgentTask(
            task_id=task_id,
            agent_type="presentation",
            title=f"Slides: {short_title}",
            prompt=topic_or_content[:100]
        )
        with self._lock:
            self.tasks[task_id] = task

        self._sync_hud()
        thread = threading.Thread(target=self._run_presentation_worker, args=(task, title, topic_or_content, slide_count), daemon=True)
        thread.start()

        return {
            "status": "launched",
            "task_id": task_id,
            "agent": "Swan Keynote Architect",
            "message": f"Presentation slide generation started for '{title}' ({slide_count} slides). PowerPoint (.pptx) and preview deck will appear on Desktop."
        }

    def _play_chime(self):
        """Plays subtle macOS completion chime."""
        try:
            sound_path = "/System/Library/Sounds/Glass.aiff"
            if os.path.exists(sound_path):
                subprocess.Popen(["afplay", sound_path], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except Exception:
            pass

    def _check_blender_http(self) -> bool:
        try:
            req = urllib.request.Request("http://127.0.0.1:9877/ping", headers={"User-Agent": "Swan/1.1"})
            with urllib.request.urlopen(req, timeout=1.0) as resp:
                if resp.status == 200:
                    return True
        except Exception:
            pass
        return False

    def _check_blender_tcp(self) -> bool:
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            s.settimeout(1.0)
            s.connect(("127.0.0.1", 9876))
            s.close()
            return True
        except Exception:
            pass
        return False

    def _ensure_blender_running(self) -> bool:
        """Checks if Blender HTTP bridge (9877) or socket (9876) is listening, or launches Blender app with bridge."""
        if self._check_blender_http() or self._check_blender_tcp():
            return True

        # Ensure startup bridge script is in place
        startup_dir = os.path.expanduser("~/Library/Application Support/Blender/5.2/scripts/startup")
        os.makedirs(startup_dir, exist_ok=True)
        bridge_dest = os.path.join(startup_dir, "blender_live_bridge.py")
        src_bridge = "/Users/ahmetyadgarov/blender/blender_live_bridge.py"
        if not os.path.isfile(bridge_dest) and os.path.isfile(src_bridge):
            try:
                import shutil
                shutil.copy2(src_bridge, bridge_dest)
            except Exception:
                pass

        try:
            print("🎨 [AgentManager] Blender not connected on port 9877/9876. Launching Blender with Live Bridge...", flush=True)
            blender_app = "/Applications/Blender.app/Contents/MacOS/Blender"
            bridge_script = bridge_dest if os.path.isfile(bridge_dest) else src_bridge
            if os.path.isfile(blender_app) and os.path.isfile(bridge_script):
                subprocess.Popen([blender_app, "--python", bridge_script], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            else:
                subprocess.Popen(["open", "-a", "Blender"])

            # Wait up to 12 seconds for bridge to be ready
            for _ in range(24):
                time.sleep(0.5)
                if self._check_blender_http() or self._check_blender_tcp():
                    print("✅ [AgentManager] Connected to Blender Live Bridge!", flush=True)
                    return True
        except Exception as e:
            print(f"⚠️ [AgentManager] Failed to launch Blender: {e}", flush=True)

        return False

    @staticmethod
    def _clean_code(raw_text: str) -> str:
        if "```python" in raw_text:
            return raw_text.split("```python")[1].split("```")[0].strip()
        elif "```py" in raw_text:
            return raw_text.split("```py")[1].split("```")[0].strip()
        elif "```" in raw_text:
            return raw_text.split("```")[1].split("```")[0].strip()
        return raw_text.strip()

    def _execute_in_blender_socket(self, code_str: str, timeout: float = 60.0) -> Dict[str, Any]:
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.settimeout(timeout)
        s.connect(("127.0.0.1", 9876))

        # Dynamically load cinematic camera rig helper if present
        rig_helper = ""
        try:
            from resource_helper import get_resource_path
            rig_file = get_resource_path("cinematic_camera_rig.py")
        except Exception:
            rig_file = os.path.join(os.path.dirname(__file__), "cinematic_camera_rig.py")
        if os.path.isfile(rig_file):
            try:
                with open(rig_file, "r", encoding="utf-8") as rf:
                    rig_helper = rf.read() + "\n"
            except Exception:
                pass

        wrapped_code = (
            "import bpy, math\n"
            "result = {'status': 'completed'}\n"
            + rig_helper
            + code_str
            + "\n"
            "# Tag 3D viewport redraw\n"
            "for wm in bpy.data.window_managers:\n"
            "    for win in wm.windows:\n"
            "        for area in win.screen.areas:\n"
            "            if area.type == 'VIEW_3D':\n"
            "                area.tag_redraw()\n"
            "# Ensure result variable is always a valid dict for Blender MCP bridge\n"
            "if not isinstance(locals().get('result'), dict):\n"
            "    result = {'status': 'completed', 'objects_count': len(bpy.data.objects)}\n"
        )

        payload = json.dumps({
            "type": "execute",
            "code": wrapped_code,
            "strict_json": False
        }) + "\0"

        s.sendall(payload.encode("utf-8"))

        data = b""
        while True:
            chunk = s.recv(8192)
            if not chunk:
                break
            data += chunk
            if b"\0" in chunk:
                break
        s.close()

        raw_response = data.decode("utf-8").rstrip("\0").strip()
        res = json.loads(raw_response)
        if isinstance(res, dict) and res.get("status") == "error":
            err_msg = res.get("message") or "Blender execution error"
            raise RuntimeError(err_msg)
        return res

    def _execute_in_blender(self, code_str: str, timeout: float = 60.0) -> Dict[str, Any]:
        """Executes python code in Blender via HTTP bridge (port 9877) or TCP socket (port 9876)."""
        # Dynamically load cinematic camera rig helper if present
        rig_helper = ""
        try:
            from resource_helper import get_resource_path
            rig_file = get_resource_path("cinematic_camera_rig.py")
        except Exception:
            rig_file = os.path.join(os.path.dirname(__file__), "cinematic_camera_rig.py")
        if os.path.isfile(rig_file):
            try:
                with open(rig_file, "r", encoding="utf-8") as rf:
                    rig_helper = rf.read() + "\n"
            except Exception:
                pass

        wrapped_code = (
            "import bpy, math\n"
            "result = {'status': 'completed'}\n"
            + rig_helper
            + code_str
            + "\n"
            "# Tag 3D viewport redraw\n"
            "for wm in bpy.data.window_managers:\n"
            "    for win in wm.windows:\n"
            "        for area in win.screen.areas:\n"
            "            if area.type == 'VIEW_3D':\n"
            "                area.tag_redraw()\n"
            "# Ensure result variable is always a valid dict for Blender MCP bridge\n"
            "if not isinstance(locals().get('result'), dict):\n"
            "    result = {'status': 'completed', 'objects_count': len(bpy.data.objects)}\n"
        )

        # 1. Try HTTP Bridge on 9877 first (runs safely on main thread via bpy.app.timers)
        if self._check_blender_http():
            try:
                req = urllib.request.Request(
                    "http://127.0.0.1:9877/execute",
                    data=json.dumps({"code": wrapped_code}).encode("utf-8"),
                    headers={"Content-Type": "application/json", "User-Agent": "Swan/1.1"}
                )
                with urllib.request.urlopen(req, timeout=timeout) as resp:
                    raw_res = resp.read().decode("utf-8")
                    res = json.loads(raw_res)
                    if not res.get("success", False):
                        err_msg = res.get("error") or res.get("message") or "Blender execution error"
                        raise RuntimeError(err_msg)
                    return {"status": "completed", "result": res}
            except Exception as http_err:
                print(f"⚠️ [AgentManager] HTTP bridge execution failed: {http_err}. Trying socket...", flush=True)

        # 2. Fallback to raw TCP socket on 9876
        return self._execute_in_blender_socket(code_str, timeout=timeout)

    @staticmethod
    def _generate_with_fallback(client, contents, config, models=("gemini-3.1-flash-lite-preview", "gemini-3.1-flash-lite", "gemini-flash-latest", "gemini-flash-lite-latest", "gemini-3.6-flash")):
        last_err = None
        from google.genai import types
        if config and not getattr(config, "http_options", None):
            config.http_options = types.HttpOptions(retry_options=types.HttpRetryOptions(attempts=1))

        for m in models:
            try:
                return client.models.generate_content(model=m, contents=contents, config=config)
            except Exception as e:
                last_err = e
                err_str = str(e)
                if any(kw in err_str for kw in ("429", "RESOURCE_EXHAUSTED", "404", "NOT_FOUND", "503", "UNAVAILABLE", "Unavailable")):
                    print(f"⚠️ [AgentManager] Model '{m}' error ({err_str[:60]}...). Falling back to next model...", flush=True)
                    continue
                raise e
        raise last_err

    def _inspect_blender_scene(self) -> Dict[str, Any]:
        """Queries Blender socket to retrieve live scene metadata before generating code."""
        try:
            inspect_code = (
                "import bpy\n"
                "objs = []\n"
                "for o in bpy.data.objects:\n"
                "    objs.append({'name': o.name, 'type': o.type, 'loc': [round(v, 2) for v in o.location]})\n"
                "active_cam = bpy.context.scene.camera.name if bpy.context.scene.camera else None\n"
                "result = {\n"
                "    'count': len(objs),\n"
                "    'objects': objs,\n"
                "    'active_camera': active_cam,\n"
                "    'frame_start': bpy.context.scene.frame_start,\n"
                "    'frame_end': bpy.context.scene.frame_end,\n"
                "    'fps': bpy.context.scene.render.fps\n"
                "}\n"
            )
            res = self._execute_in_blender(inspect_code, timeout=8.0)
            return res.get("result", {})
        except Exception as e:
            print(f"⚠️ [AgentManager] Scene inspection failed: {e}", flush=True)
            return {}

    def _run_blender_worker(self, task: AgentTask, style: str, reference_image_path: Optional[str] = None):
        print(f"🎨 [Blender Agent] Processing 3D task: '{task.prompt}' (reference: {reference_image_path})", flush=True)
        try:
            from google import genai
            from google.genai import types
            client = genai.Client(
                api_key=config.api_key,
                http_options={"retry_options": {"attempts": 1}}
            )

            if not self._ensure_blender_running():
                raise RuntimeError("Could not connect to Blender on port 9877 or 9876. Please ensure Blender is running.")

            # MODE A: 2D REFERENCE IMAGE / LOGO TO 3D MODEL RECONSTRUCTION
            if reference_image_path and os.path.isfile(reference_image_path):
                from PIL import Image
                pil_img = Image.open(reference_image_path)
                img_name = os.path.basename(reference_image_path)
                print(f"👁️ [Blender Agent] Mode: 2D REFERENCE LOGO RECONSTRUCTION for '{img_name}'", flush=True)
                _update_hud_working("Agent vision active 👁️", f"Analyzing {img_name}...")

                system_instruction = (
                    "You are a World-Class 3D Technical Artist, Master Sculptor, and Blender 5.2 Python (bpy) Automation Specialist.\n"
                    "Your mission is to analyze this 2D reference logo / graphic image and write a complete, self-contained Blender 5.2 Python script that creates a stunning, TRUE 3D VOLUMETRIC SCULPTURAL MASTERPIECE in Blender.\n\n"
                    "CRITICAL VOLUMETRIC 3D SCULPTING RULES (NEVER CREATE FLAT 2D EXTRUSIONS):\n"
                    "The user wants a genuine 3D sculptural model, NOT a flat silhouette extruded on a 2D plane!\n\n"
                    "1. CLEAN SLATE & PURGE:\n"
                    "   for o in list(bpy.data.objects): bpy.data.objects.remove(o, do_unlink=True)\n"
                    "   for m in list(bpy.data.materials): bpy.data.materials.remove(m, do_unlink=True)\n"
                    "   for c in list(bpy.data.curves): bpy.data.curves.remove(c, do_unlink=True)\n"
                    "   for img in list(bpy.data.images): bpy.data.images.remove(img, do_unlink=True)\n\n"
                    "2. CIRCLES, RINGS & DONUTS MUST BE TRUE 3D TORUSES:\n"
                    "   - Any circular elements, rings, or loops (like top and bottom circles in logos) MUST be modeled as true 3D Toruses:\n"
                    "     bpy.ops.mesh.primitive_torus_add(\n"
                    "         major_radius=0.40, minor_radius=0.09,\n"
                    "         major_segments=64, minor_segments=32,\n"
                    "         location=(0.0, 0.0, z_pos),\n"
                    "         rotation=(math.radians(90), 0, 0)\n"
                    "     )\n"
                    "     obj = bpy.context.active_object\n"
                    "     bpy.ops.object.shade_smooth()\n\n"
                    "3. WINGS, CRESTS & OUTLINES MUST BE CLOSED 3D VOLUMETRIC CONTOUR LOOPS (NEVER SINGLE OPEN STICKS):\n"
                    "   - The wings in logos are CLOSED RIBBONS / CONTOUR LOOPS enclosing negative space! NEVER make an open single line:\n"
                    "     c = bpy.data.curves.new(name='Wing_Curve', type='CURVE')\n"
                    "     c.dimensions = '3D'\n"
                    "     c.bevel_depth = 0.085  # thick, lustrous rounded 3D volume\n"
                    "     c.bevel_resolution = 6\n"
                    "     c.use_fill_caps = True\n"
                    "     spline = c.splines.new(type='BEZIER')\n"
                    "     spline.use_cyclic_u = True  # MUST be closed loop!\n"
                    "   - The spline MUST trace the COMPLETE closed loop: root V-notch -> upper arch -> wingtip -> under-wing return -> back to root!\n"
                    "     For example for upper wing:\n"
                    "       (0.01, 0.0, 0.00) [root V-notch, handle VECTOR],\n"
                    "       (0.40, -0.06, 0.36) [arching up],\n"
                    "       (1.20, -0.16, 0.58) [high crest],\n"
                    "       (2.40, -0.30, 0.54) [long sweep],\n"
                    "       (3.15, -0.42, 0.42) [outer wingtip, handle VECTOR],\n"
                    "       (2.60, -0.32, 0.22) [under-tip tuck],\n"
                    "       (1.70, -0.18, 0.22) [lower contour],\n"
                    "       (0.80, -0.08, 0.18) [inner return]\n"
                    "     And similarly for lower wing: a complete closed loop from inner root (0.45, -0.05, -0.18) swooping out to tip (2.80, -0.38, -0.02) and returning via belly (0.60, -0.06, -0.45)!\n"
                    "   - AERODYNAMIC 3D CURVATURE IN 3D SPACE:\n"
                    "     Points arch forward along Y as they extend outward in X, giving genuine 3D aeronautical depth!\n"
                    "   - SYMMETRY VIA MIRROR MODIFIER:\n"
                    "     Model the right side accurately, then add Mirror Modifier: mir = obj.modifiers.new('Mirror', 'MIRROR'); mir.use_axis[0] = True.\n\n"
                    "4. PBR TEAL / MULTI-TONE LACQUER MATERIAL WITH CLEARCOAT:\n"
                    "   - Extract exact colors from logo image. Convert sRGB to Linear RGB.\n"
                    "   - Principled BSDF setup:\n"
                    "     'Base Color': (0.010, 0.44, 0.41, 1.0) for teal or logo linear color.\n"
                    "     'Roughness': 0.14 (luxurious satin gloss).\n"
                    "     'Metallic': 0.20 (subtle metallic flake depth).\n"
                    "     'Coat Weight': 0.85 (shiny automotive clearcoat lacquer).\n"
                    "     'Coat Roughness': 0.03.\n\n"
                    "5. STUDIO ENVIRONMENT & REFLECTIVE PEDESTAL:\n"
                    "   - Add dark reflective floor pedestal below model catching soft drop shadows:\n"
                    "     bpy.ops.mesh.primitive_cylinder_add(radius=7.0, depth=0.25, vertices=64, location=(0, 0, -2.0))\n"
                    "     Floor material: Base Color (0.012, 0.012, 0.018, 1.0), Roughness 0.18, Metallic 0.5.\n\n"
                    "6. THREE-POINT HOLLYWOOD STUDIO LIGHTING RIG:\n"
                    "   - Key Light: AREA light, Energy 800W, size 3.5, warm tone (1.0, 0.98, 0.95), location (5.0, -6.0, 4.0).\n"
                    "   - Fill Light: AREA light, Energy 350W, size 5.0, soft blue-cyan tone (0.80, 0.92, 1.0), location (-6.0, -5.0, 3.0).\n"
                    "   - Rim Light: AREA light, Energy 950W, size 4.0, cyan-white tone (0.55, 0.95, 1.0), location (0.0, 6.0, 3.5) catching the 3D beveled edges.\n\n"
                    "7. DYNAMIC 3D CAMERA PRESENTATION & 360-DEGREE TURNTABLE ORBIT (NEVER FLAT FRONT VIEW):\n"
                    "   - Create Camera_LookTarget Empty at subject center (0.0, -0.15, 0.05).\n"
                    "   - Camera with TRACK_TO constraint pointing at Camera_LookTarget. bpy.context.scene.camera = cam_obj.\n"
                    "   - Focal lens: 45mm, Depth of Field enabled (cam.data.dof.use_dof = True, cam.data.dof.focus_object = cam_target, aperture_fstop = 2.8).\n"
                    "   - Position at dynamic 3/4 perspective elevation angle (~35 deg elevation, ~35 deg azimuth, distance r ~ 9.0m) so the volumetric 3D thickness is immediately apparent.\n"
                    "   - Master Timeline: scene.render.fps = 24, scene.frame_start = 1, scene.frame_end = 120.\n"
                    "   - Animate smooth Hollywood turntable orbit around the 3D sculpture from frame 1 to 120:\n"
                    "     r = 9.2; z_base = 2.8\n"
                    "     for frame in range(1, 121):\n"
                    "         t = (frame - 1) / 120.0\n"
                    "         angle = math.radians(35) + t * 2 * math.pi\n"
                    "         cam.location = (math.sin(angle) * r, -math.cos(angle) * r, z_base + 0.5 * math.sin(t * 2 * math.pi))\n"
                    "         cam.keyframe_insert(data_path='location', frame=frame)\n"
                    "     scene.frame_set(1)\n\n"
                    "8. Output ONLY raw executable Python code inside a ```python ``` markdown codeblock without explanations."
                )

                user_msg = f"Recreate this reference logo ({img_name}) as a stunning 3D model in Blender 5.2.\nUser request: {task.prompt}"
                cfg = types.GenerateContentConfig(
                    system_instruction=system_instruction,
                    temperature=0.2,
                    http_options=types.HttpOptions(retry_options=types.HttpRetryOptions(attempts=1))
                )

                _update_hud_working("Modeling in Blender 3D...", f"Recreating 3D {img_name}...")
                response = self._generate_with_fallback(client, [pil_img, user_msg], cfg)
                code = self._clean_code(response.text or "")
                is_modification = False
            else:
                # 1. Live scene introspection
                scene_info = self._inspect_blender_scene()
                existing_count = scene_info.get("count", 0)
                existing_objs = scene_info.get("objects", [])
                active_cam = scene_info.get("active_camera")

                prompt_lower = task.prompt.lower()
                clean_slate_triggers = [
                    "new scene", "brand new", "start over", "clean slate", "clear scene",
                    "clear everything", "yangi sahna", "boshidan boshla", "tozala"
                ]
                is_clean_slate_requested = any(trig in prompt_lower for trig in clean_slate_triggers)
                is_modification = (existing_count > 0) and not is_clean_slate_requested

                if is_modification:
                    obj_summary = ", ".join([f"{o['name']} ({o['type']})" for o in existing_objs[:15]])
                    if len(existing_objs) > 15:
                        obj_summary += f"... (+{len(existing_objs)-15} more)"

                    print(f"🎬 [Blender Agent] Mode: SCENE PRESERVATION ({existing_count} existing objects). Updating camera/scene without deleting meshes.", flush=True)

                    system_instruction = (
                        "You are an elite Hollywood 3D Director and master Blender Python developer for Blender 5.2.\n"
                        "You are working on an EXISTING Blender 3D scene with existing objects the user loves.\n\n"
                        "CRITICAL SCENE PRESERVATION (ABSOLUTE NON-NEGOTIABLE RULE):\n"
                        "- NEVER delete, remove, unlink, or modify existing meshes, curves, or materials! DO NOT run bpy.data.objects.remove().\n"
                        "- Retain ALL existing objects and scene geometry in place.\n"
                        "- The user wants camera choreography, camera movement, or scene animation.\n\n"
                        "HOLLYWOOD CINEMATOGRAPHY ENGINE & 4 ARCHETYPES:\n"
                        "Map the user's request to the appropriate cinematography archetype:\n"
                        "1. ARCHETYPE 1: CONTINUOUS ONE-SHOT PREVIZ & LIVING LENS (for general flow, orbits, tracking):\n"
                        "   - Single continuous camera move along a smooth organic spline.\n"
                        "   - Living Lens: Animate focal length (cam.data.keyframe_insert(data_path='lens', frame=f)). Widen to 22-24mm during fast movement, tighten to 35-50mm on pauses/accents.\n"
                        "   - Smooth trigonometric orbital path: cx = hero_x + cos(angle)*radius, cy = hero_y + sin(angle)*radius, cz = hero_z + height + sin(t*pi)*tilt.\n"
                        "2. ARCHETYPE 2: DIALOGUE RAILS & OTS COVERAGE (for conversations, multi-character):\n"
                        "   - 21:9 cinematic framing (resolution_x=1920, resolution_y=804).\n"
                        "   - Over-the-shoulder (OTS) shots crawling at shoulder height.\n"
                        "   - Instant cut discipline: camera location and target jump instantly on cut frames with zero transition drift.\n"
                        "3. ARCHETYPE 3: HIGH-CONCEPT ACTION, SURFACE DIVES & ROBO-ARM:\n"
                        "   - Vertical elevator / floor dive: PURE Z-axis movement ONLY, zero rotation (lock X and Y coordinates), smooth ease-in, slowdown in middle with hero dead center.\n"
                        "   - Robo-Arm snaps: Camera whip-arcs to distinct angle, followed by a DEAD STOP with ZERO drift (keyframe identical coordinates at hold start and hold end).\n"
                        "4. ARCHETYPE 4: HYPERMOTION COMMERCIAL:\n"
                        "   - Speed ramping: every shot snaps in hard, sags in the middle, and accelerates into the cut.\n"
                        "   - Packshot composition: hero subject framed on the left third, breathing room on the right.\n\n"
                        "CAMERA RIGGING ARCHITECTURE (BLENDER 5.2 EXACT SYNTAX):\n"
                        "- Identify hero object/center of interest from existing objects.\n"
                        "- Clean old animation on camera: if cam.animation_data: cam.animation_data_clear(). If cam.data.animation_data: cam.data.animation_data_clear().\n"
                        "- Create or locate Empty object named `Camera_LookTarget` placed at subject center: (scene.collection.objects.link(look_target)).\n"
                        "- Ensure TRACK_TO constraint on camera: track_axis='TRACK_NEGATIVE_Z', up_axis='UP_Y', target=look_target.\n"
                        "- Blender 5.2 DOF: cam.data.dof.use_dof = True; cam.data.dof.focus_object = look_target; cam.data.dof.aperture_fstop = 2.8 (NEVER use focus_target).\n"
                        "- Anti-Gimbal Singularity: When passing overhead, maintain minimum horizontal offset (abs(x) >= 0.04 or abs(y) >= 0.04) to prevent 180-degree flip.\n"
                        "- Standardize timeline: scene.render.fps = 24; scene.frame_start = 1; scene.frame_end = 120 (or matched duration); scene.frame_set(1).\n"
                        "- Render engine: Leave existing engine intact.\n"
                        "- NEVER use bpy.ops.wm.read_factory_settings, bpy.ops.wm.read_homefile, or sys.exit.\n"
                        "- Output ONLY executable Python code inside a ```python ``` markdown codeblock without explanations."
                    )

                    user_msg = (
                        f"CURRENT LIVE BLENDER SCENE:\n"
                        f"- Existing Objects ({existing_count}): {obj_summary}\n"
                        f"- Active Camera: {active_cam or 'None'}\n\n"
                        f"USER REQUEST: {task.prompt}\n"
                        f"Cinematography Style: {style}\n"
                        f"TASK: Update camera choreography and scene elements smoothly. DO NOT delete existing objects!"
                    )
                else:
                    print(f"🎬 [Blender Agent] Mode: CLEAN SLATE / NEW SCENE. Generating complete 3D scene.", flush=True)

                    system_instruction = (
                        "You are an elite Hollywood 3D Director and master Blender Python developer for Blender 5.2.\n"
                        "The user is creating a BRAND NEW 3D scene from scratch.\n\n"
                        "CRITICAL BLENDER 5.2 RULES (VIOLATIONS WILL CRASH THE ENGINE):\n"
                        "1. Clean slate: remove existing objects:\n"
                        "   for obj in list(bpy.data.objects): bpy.data.objects.remove(obj, do_unlink=True)\n"
                        "   for mat in list(bpy.data.materials): bpy.data.materials.remove(mat, do_unlink=True)\n"
                        "2. Mesh Primitive Operators (exact names & parameters):\n"
                        "   - bpy.ops.mesh.primitive_cube_add(size=..., location=...)\n"
                        "   - bpy.ops.mesh.primitive_cylinder_add(radius=..., depth=..., vertices=..., location=...)\n"
                        "   - bpy.ops.mesh.primitive_plane_add(size=..., location=...)\n"
                        "   - bpy.ops.mesh.primitive_uv_sphere_add(radius=..., segments=..., ring_count=..., location=...)\n"
                        "   - bpy.ops.mesh.primitive_ico_sphere_add(radius=..., subdivisions=..., location=...) (NOTE: 'ico_sphere' with underscore)\n"
                        "   - bpy.ops.mesh.primitive_torus_add(major_radius=..., minor_radius=..., location=...)\n"
                        "   - bpy.ops.mesh.primitive_cone_add(radius1=..., depth=..., location=...)\n"
                        "3. In Blender 5.x / 4.x, Principled BSDF input sockets are: 'Base Color', 'Metallic', 'Roughness', 'Emission Color', 'Emission Strength', 'Specular IOR Level'.\n"
                        "4. Camera & Cinematic Staging:\n"
                        "   - Create a Camera_LookTarget Empty at subject position.\n"
                        "   - Add TRACK_TO constraint on camera pointing at LookTarget.\n"
                        "   - Depth of field: cam.data.dof.focus_object = look_target.\n"
                        "   - Animate dynamic camera movement: cam.keyframe_insert(data_path='location', frame=f).\n"
                        "   - Master timeline: scene.frame_start = 1, scene.frame_end = 120, scene.frame_set(1).\n"
                        "5. Output ONLY raw executable Python code inside a ```python ``` markdown codeblock. No commentary."
                    )

                    user_msg = f"Build this brand new 3D scene in Blender with {style} cinematography:\nPrompt: {task.prompt}"

                cfg = types.GenerateContentConfig(
                    system_instruction=system_instruction,
                    temperature=0.35,
                    http_options=types.HttpOptions(retry_options=types.HttpRetryOptions(attempts=1))
                )

                response = self._generate_with_fallback(client, user_msg, cfg)
                code = self._clean_code(response.text or "")

            # Send code to Blender with self-healing retry (up to 2 repair attempts)
            current_code = code
            last_err = None
            for attempt in range(3):
                try:
                    exec_res = self._execute_in_blender(current_code)
                    print(f"🎨 [Blender Agent] Code executed successfully (attempt {attempt+1}): {exec_res}", flush=True)
                    last_err = None
                    break
                except RuntimeError as exec_err:
                    last_err = exec_err
                    print(f"⚠️ [Blender Agent] Attempt {attempt+1} failed: {exec_err}", flush=True)
                    if attempt < 2:
                        _update_hud_working("Agent repairing...", f"Self-healing syntax (try {attempt+1})...")
                        repair_prompt = (
                            f"The following Blender Python script failed with this runtime error in Blender 5.2:\n"
                            f"ERROR: {exec_err}\n\n"
                            f"FAILED CODE:\n{current_code}\n\n"
                            f"Please fix the error and output ONLY the corrected complete executable Python code inside a ```python ``` markdown block.\n"
                            f"IMPORTANT RULES FOR BLENDER 5.2:\n"
                            f"- DO NOT delete existing objects if in preservation mode.\n"
                            f"- Camera DOF: use cam.data.dof.focus_object = look_target (NOT focus_target).\n"
                            f"- Keyframing: use cam.keyframe_insert(data_path='location', frame=f). Do NOT access action.fcurves directly.\n"
                            f"- Leave render engine untouched.\n"
                        )
                        repair_resp = self._generate_with_fallback(client, repair_prompt, cfg)
                        current_code = self._clean_code(repair_resp.text or "")
                    else:
                        raise last_err

            if last_err:
                raise last_err

            # Verification: Ensure objects were actually created or preserved
            verify_res = self._execute_in_blender("result = {'count': len(bpy.data.objects), 'names': [o.name for o in bpy.data.objects]}")
            obj_info = verify_res.get("result", {})
            obj_count = obj_info.get("count", 0)

            if obj_count == 0:
                raise RuntimeError("Blender execution finished, but 0 3D objects exist in the scene.")

            # Bring Blender to focus
            try:
                subprocess.run(["osascript", "-e", 'tell application "Blender" to activate'], capture_output=True)
            except Exception:
                pass

            if reference_image_path and os.path.isfile(reference_image_path):
                img_bname = os.path.basename(reference_image_path)
                action_desc = f"3D model of '{img_bname}' created"
                hud_msg = f"3D {img_bname} ready"
            elif is_modification:
                action_desc = "Camera & scene choreography updated"
                hud_msg = f"Blender: {task.prompt[:25]}"
            else:
                action_desc = "3D scene created"
                hud_msg = f"Blender: {task.prompt[:25]}"

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = f"{action_desc} successfully with {obj_count} objects in Blender! Press Spacebar in Blender to play animation."

            # Synchronize multi-agent HUD
            self._sync_hud(recent_title=hud_msg, is_completion=True)
            self._play_chime()

        except Exception as e:
            print(f"❌ [Blender Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="Blender failed", is_completion=True)

    def _run_generic_worker(self, task: AgentTask, details: str):
        print(f"🤖 [Generic Agent] Starting task: '{task.title}'", flush=True)
        try:
            from google import genai
            from google.genai import types
            client = genai.Client(api_key=config.api_key)

            prompt = f"Perform this autonomous task thoroughly:\nTask: {task.prompt}\nDetails: {details}"
            cfg = types.GenerateContentConfig(temperature=0.3)
            response = self._generate_with_fallback(client, prompt, cfg)

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = response.text or "Task completed"

            self._sync_hud(recent_title=task.title[:25], is_completion=True)
            self._play_chime()
        except Exception as e:
            print(f"❌ [Generic Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="Agent failed", is_completion=True)

    def _run_image_worker(self, task: AgentTask, source_image_path: Optional[str], aspect_ratio: str, model_preference: Optional[str]):
        print(f"🎨 [Image Agent] Starting image task: '{task.title}' (aspect: {aspect_ratio})", flush=True)
        try:
            from datetime import datetime
            desktop_dir = os.path.expanduser("~/Desktop")
            os.makedirs(desktop_dir, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

            is_edit = bool(source_image_path)
            if is_edit:
                dest_filename = f"Swan_Edited_{timestamp}.png"
            else:
                dest_filename = f"Swan_Image_{timestamp}.png"
            dest_path = os.path.join(desktop_dir, dest_filename)

            # Resolve source image if editing
            pil_source = None
            resolved_source_path = None
            if is_edit and source_image_path:
                raw_path = os.path.expanduser(source_image_path.strip())
                if os.path.isfile(raw_path):
                    resolved_source_path = raw_path
                elif os.path.isdir(raw_path):
                    exts = (".png", ".jpg", ".jpeg", ".webp", ".bmp")
                    imgs = [os.path.join(raw_path, f) for f in os.listdir(raw_path) if f.lower().endswith(exts)]
                    if imgs:
                        imgs.sort(key=lambda p: os.path.getmtime(p), reverse=True)
                        resolved_source_path = imgs[0]
                else:
                    for folder in [desktop_dir, os.path.expanduser("~/Downloads"), os.path.expanduser("~/Pictures")]:
                        candidate = os.path.join(folder, raw_path.lstrip("/~"))
                        if os.path.isfile(candidate):
                            resolved_source_path = candidate
                            break
                        if raw_path.lower() in folder.lower() and os.path.isdir(folder):
                            exts = (".png", ".jpg", ".jpeg", ".webp")
                            imgs = [os.path.join(folder, f) for f in os.listdir(folder) if f.lower().endswith(exts)]
                            if imgs:
                                imgs.sort(key=lambda p: os.path.getmtime(p), reverse=True)
                                resolved_source_path = imgs[0]
                                break

                if resolved_source_path and os.path.isfile(resolved_source_path):
                    print(f"🖼️ [Image Agent] Loaded source image for editing: {resolved_source_path}", flush=True)
                    from PIL import Image
                    pil_source = Image.open(resolved_source_path)
                else:
                    print(f"⚠️ [Image Agent] Could not locate source image '{source_image_path}'. Proceeding with descriptive generation.", flush=True)

            dim_map = {
                "1:1": (1024, 1024),
                "16:9": (1280, 720),
                "9:16": (720, 1280),
                "4:3": (1024, 768),
                "3:4": (768, 1024)
            }
            norm_ratio = aspect_ratio if aspect_ratio in dim_map else "1:1"
            width, height = dim_map[norm_ratio]

            candidate_models = ["gemini-2.5-flash-image", "gemini-3.1-flash-image", "gemini-3-pro-image", "gemini-3.1-flash-lite-image"]
            if model_preference:
                pref = model_preference.lower()
                if "lite" in pref:
                    candidate_models = ["gemini-3.1-flash-lite-image", "gemini-2.5-flash-image"]
                elif "pro" in pref:
                    candidate_models = ["gemini-3-pro-image", "gemini-3.1-flash-image"]

            success = False
            try:
                from google import genai
                from google.genai import types
                client = genai.Client(api_key=config.api_key)

                contents = [pil_source, task.prompt] if pil_source else task.prompt
                image_cfg = types.ImageConfig(aspect_ratio=norm_ratio)
                gen_cfg = types.GenerateContentConfig(
                    response_modalities=["IMAGE"],
                    image_config=image_cfg
                )

                for m in candidate_models:
                    try:
                        print(f"🎨 [Image Agent] Trying Google GenAI model: '{m}'...", flush=True)
                        resp = client.models.generate_content(model=m, contents=contents, config=gen_cfg)
                        if resp.candidates and resp.candidates[0].content and resp.candidates[0].content.parts:
                            for part in resp.candidates[0].content.parts:
                                if getattr(part, "inline_data", None) and part.inline_data.data:
                                    with open(dest_path, "wb") as f:
                                        f.write(part.inline_data.data)
                                    success = True
                                    print(f"✅ [Image Agent] Successfully generated image via Google GenAI ({m})!", flush=True)
                                    break
                                elif hasattr(part, "as_image"):
                                    part.as_image().save(dest_path, "PNG")
                                    success = True
                                    print(f"✅ [Image Agent] Successfully saved image via Google GenAI ({m})!", flush=True)
                                    break
                        if success:
                            break
                    except Exception as model_err:
                        print(f"⚠️ [Image Agent] Model '{m}' returned: {model_err}", flush=True)
                        continue
            except Exception as genai_err:
                print(f"⚠️ [Image Agent] Google GenAI setup notice: {genai_err}", flush=True)

            if not success or not os.path.exists(dest_path) or os.path.getsize(dest_path) < 1000:
                print("🎨 [Image Agent] Utilizing robust high-resolution visual engine with watermark purge...", flush=True)
                final_prompt = task.prompt
                if pil_source:
                    try:
                        from google import genai
                        client = genai.Client(api_key=config.api_key)
                        v_resp = client.models.generate_content(
                            model="gemini-flash-latest",
                            contents=[pil_source, f"Describe this image and incorporate this instruction: '{task.prompt}'. Output a single photorealistic generation prompt describing the modified scene."]
                        )
                        if v_resp.text:
                            final_prompt = v_resp.text.strip().replace("\n", " ")
                    except Exception:
                        pass
                else:
                    try:
                        from google import genai
                        client = genai.Client(api_key=config.api_key)
                        enh_resp = client.models.generate_content(
                            model="gemini-flash-latest",
                            contents=(
                                "You are an award-winning master digital visual artist.\n"
                                f"Convert this image request into an ultra-detailed, photorealistic, cinematic visual description: '{task.prompt}'.\n"
                                "Specify: subject details, materials, studio lighting, octane render, photorealistic, 8k resolution, cinematic composition.\n"
                                "Output ONLY the prompt string in English, no quotes, concise within 250 characters."
                            )
                        )
                        if enh_resp.text:
                            final_prompt = enh_resp.text.strip().replace("\n", " ")
                    except Exception:
                        pass

                import random
                ctx = ssl.create_default_context()
                ctx.check_hostname = False
                ctx.verify_mode = ssl.CERT_NONE

                clean_prompt = final_prompt[:250].strip()
                enc_p = urllib.parse.quote(clean_prompt)
                seed = random.randint(1000, 999999)
                raw_bytes = b""

                urls_to_try = [
                    f"https://image.pollinations.ai/prompt/{enc_p}?width={width}&height={height}&model=flux&nologo=true&seed={seed}",
                    f"https://image.pollinations.ai/prompt/{enc_p}?width={width}&height={height}&model=turbo&nologo=true&seed={seed}",
                    f"https://image.pollinations.ai/prompt/{enc_p}?width={width}&height={height}&nologo=true&seed={seed}",
                    f"https://image.pollinations.ai/prompt/{urllib.parse.quote(task.prompt[:180].strip())}?width={width}&height={height}&nologo=true&seed={seed}"
                ]

                for u in urls_to_try:
                    try:
                        req_p = urllib.request.Request(u, headers={"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)"})
                        with urllib.request.urlopen(req_p, timeout=25, context=ctx) as resp:
                            data = resp.read()
                            if len(data) > 1000:
                                raw_bytes = data
                                break
                    except Exception as net_err:
                        print(f"⚠️ [Image Agent] URL attempt failed ({u[:60]}...): {net_err}", flush=True)

                if len(raw_bytes) > 1000:
                    from PIL import Image
                    import io
                    raw_img = Image.open(io.BytesIO(raw_bytes))
                    w, h = raw_img.size
                    clean_img = raw_img.crop((0, 0, w, max(10, h - 48)))
                    clean_img = clean_img.resize((width, height), Image.Resampling.LANCZOS)
                    clean_img.save(dest_path, "PNG", quality=95)
                    success = True

            if not os.path.exists(dest_path) or os.path.getsize(dest_path) < 1000:
                raise RuntimeError("Failed to generate and save image file to Desktop.")

            try:
                subprocess.Popen(["open", dest_path])
            except Exception:
                pass

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = f"Image saved to Desktop: {dest_path}"

            # Synchronize multi-agent HUD
            self._sync_hud(recent_title=dest_filename, is_completion=True)
            self._play_chime()

        except Exception as e:
            print(f"❌ [Image Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="Image failed", is_completion=True)

    def _run_transcribe_worker(self, task: AgentTask, file_path: Optional[str], query: Optional[str]):
        print(f"🎙️ [Transcribe Agent] Starting audio transcription for '{file_path or query}'", flush=True)
        _update_hud_working("Transcribing Audio 🎙️", "Locating audio file...")
        try:
            desktop_dir = os.path.expanduser("~/Desktop")
            search_folders = [
                os.path.expanduser("~/Downloads"),
                desktop_dir,
                os.path.expanduser("~/Documents"),
                os.path.expanduser("~/Music")
            ]
            audio_exts = (".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".webm")

            resolved_path = None
            target = (file_path or query or "").strip()

            if target:
                expanded = os.path.expanduser(target)
                if os.path.isfile(expanded):
                    resolved_path = expanded
                else:
                    clean_norm = os.path.basename(target).lower().replace(" ", "").replace("_", "").replace("-", "")
                    for folder in search_folders:
                        if not os.path.isdir(folder):
                            continue
                        for f in os.listdir(folder):
                            if f.lower().endswith(audio_exts):
                                f_norm = f.lower().replace(" ", "").replace("_", "").replace("-", "")
                                if clean_norm in f_norm or f_norm in clean_norm:
                                    resolved_path = os.path.join(folder, f)
                                    break
                        if resolved_path:
                            break

            if not resolved_path:
                all_audios = []
                for folder in search_folders:
                    if os.path.isdir(folder):
                        for f in os.listdir(folder):
                            if f.lower().endswith(audio_exts):
                                p = os.path.join(folder, f)
                                all_audios.append((os.path.getmtime(p), p))
                if all_audios:
                    all_audios.sort(key=lambda x: x[0], reverse=True)
                    resolved_path = all_audios[0][1]

            if not resolved_path or not os.path.isfile(resolved_path):
                raise RuntimeError("Could not find any audio file to transcribe in ~/Downloads, ~/Desktop, ~/Documents, or ~/Music.")

            audio_name = os.path.basename(resolved_path)
            print(f"🎙️ [Transcribe Agent] Found audio file: {resolved_path}", flush=True)
            _update_hud_working("Transcribing Audio 🎙️", f"Analyzing {audio_name[:25]}...")

            from google import genai
            from google.genai import types
            client = genai.Client(api_key=config.api_key)

            with open(resolved_path, "rb") as af:
                audio_bytes = af.read()

            ext = os.path.splitext(resolved_path)[1].lower()
            mime_map = {
                ".wav": "audio/wav",
                ".mp3": "audio/mp3",
                ".m4a": "audio/mp4",
                ".aac": "audio/aac",
                ".ogg": "audio/ogg",
                ".flac": "audio/flac",
                ".webm": "audio/webm"
            }
            mime = mime_map.get(ext, "audio/mp3")

            part = types.Part.from_bytes(data=audio_bytes, mime_type=mime)
            prompt = (
                "You are an expert, meticulous audio transcriber and speech-to-text intelligence. "
                "Please transcribe the provided audio verbatim, completely, and accurately in the exact language spoken. "
                "Format with clear speaker turns or natural paragraph breaks. "
                "Do NOT add introductory or concluding chatter. Output ONLY the clean transcription."
            )

            try:
                resp = self._generate_with_fallback(client, [part, prompt], None, models=("gemini-flash-latest", "gemini-3.1-flash-lite-preview", "gemini-flash-lite-latest"))
                transcript = resp.text.strip() if resp and resp.text else ""
            except Exception as te:
                print(f"⚠️ [Transcribe Agent] Primary audio model notice: {te}. Retrying with flash-lite...", flush=True)
                resp = client.models.generate_content(model="gemini-3.1-flash-lite-preview", contents=[part, prompt])
                transcript = resp.text.strip() if resp and resp.text else ""

            if not transcript:
                raise RuntimeError("Transcription result was empty.")

            base_root = os.path.splitext(audio_name)[0]
            out_file = os.path.join(desktop_dir, f"{base_root}_transcript.md")
            with open(out_file, "w", encoding="utf-8") as out_f:
                out_f.write(f"# Audio Transcription: {audio_name}\n\n")
                out_f.write(f"*Transcribed by Swan AI Agent on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n\n")
                out_f.write("---\n\n")
                out_f.write(transcript)
                out_f.write("\n")

            try:
                subprocess.Popen(["open", out_file])
            except Exception:
                pass

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = f"Audio transcribed successfully: {out_file}"
            self._sync_hud(recent_title=f"{base_root[:20]} transcribed", is_completion=True)
            self._play_chime()

        except Exception as e:
            print(f"❌ [Transcribe Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="Transcription failed", is_completion=True)

    def _run_youtube_worker(self, task: AgentTask, query_or_url: str, focus: str):
        print(f"▶️ [YouTube Agent] Searching & summarizing: '{query_or_url}' (focus: {focus})", flush=True)
        _update_hud_working("YouTube Agent ▶️", "Locating video & captions...")
        try:
            import yt_dlp
            import requests

            target_url = query_or_url.strip()
            video_title = "YouTube Video"
            video_url = ""

            ydl_opts = {
                "quiet": True,
                "skip_download": True,
                "no_warnings": True
            }

            is_url = "youtube.com" in target_url or "youtu.be" in target_url
            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                search_key = target_url if is_url else f"ytsearch1:{target_url}"
                info = ydl.extract_info(search_key, download=False)
                if "entries" in info and info["entries"]:
                    entry = info["entries"][0]
                else:
                    entry = info

                video_title = entry.get("title") or "YouTube Video"
                video_url = entry.get("webpage_url") or entry.get("url") or target_url
                uploader = entry.get("uploader") or entry.get("channel") or ""
                description = entry.get("description") or ""
                auto_caps = entry.get("automatic_captions") or {}
                subs = entry.get("subtitles") or {}

            print(f"▶️ [YouTube Agent] Video identified: '{video_title}' ({video_url})", flush=True)
            _update_hud_working("YouTube Agent ▶️", f"Extracting {video_title[:25]}...")

            captions_text = ""
            all_caps = {**auto_caps, **subs}
            lang_pool = ["uz", "en", "ru", "tr"]
            found_caps = None
            for lp in lang_pool:
                if lp in all_caps:
                    found_caps = all_caps[lp]
                    break
            if not found_caps and all_caps:
                found_caps = list(all_caps.values())[0]

            if found_caps:
                sub_url = None
                for fmt in found_caps:
                    if fmt.get("ext") in ("json3", "srv3", "vtt"):
                        sub_url = fmt.get("url")
                        break
                if not sub_url and found_caps:
                    sub_url = found_caps[0].get("url")

                if sub_url:
                    try:
                        resp = requests.get(sub_url, headers={"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)"}, timeout=12)
                        if "json" in resp.headers.get("Content-Type", "") or sub_url.endswith("json3") or "json" in sub_url:
                            data = resp.json()
                            lines = []
                            for ev in data.get("events", []):
                                segs = ev.get("segs", [])
                                txt = "".join(s.get("utf8", "") for s in segs).strip()
                                if txt and txt != "\n":
                                    t_sec = ev.get("tStartMs", 0) // 1000
                                    lines.append(f"[{t_sec // 60:02d}:{t_sec % 60:02d}] {txt}")
                            captions_text = "\n".join(lines)
                        else:
                            captions_text = resp.text
                    except Exception as ce:
                        print(f"⚠️ [YouTube Agent] Caption download notice: {ce}", flush=True)

            if not captions_text:
                captions_text = f"Video Description:\n{description[:4000]}"

            _update_hud_working("YouTube Agent ▶️", "Generating Executive Summary...")
            from google import genai
            from google.genai import types
            client = genai.Client(api_key=config.api_key)

            summary_prompt = (
                "You are an elite research analyst and executive summarizer.\n"
                f"Video Title: {video_title}\n"
                f"Channel: {uploader}\n"
                f"URL: {video_url}\n"
                f"User Focus / Request: {focus or 'Comprehensive high-level summary'}\n\n"
                f"Transcript / Content:\n{captions_text[:35000]}\n\n"
                "Please produce a world-class, beautifully structured Executive Summary in Markdown with:\n"
                "1. 🎯 Executive Overview & Core Thesis (2-3 concise paragraphs)\n"
                "2. 💡 Key Takeaways & Actionable Insights (bulleted)\n"
                "3. ⏱️ Timeline & Section Highlights (key timestamp moments)\n"
                "4. 💬 Notable Quotes / Standout Statements\n"
                "5. 🏁 Conclusion & Final Synthesis\n"
                "Format in crisp, modern GitHub Markdown with bolding, quotes, and clean dividers."
            )

            res = self._generate_with_fallback(
                client,
                contents=summary_prompt,
                config=None,
                models=("gemini-flash-latest", "gemini-3.1-flash-lite-preview", "gemini-flash-lite-latest", "gemini-2.5-flash")
            )
            summary_content = res.text.strip() if res and res.text else "Failed to generate summary."

            safe_title = "".join(c for c in video_title if c.isalnum() or c in (" ", "_", "-")).strip()[:40]
            desktop_dir = os.path.expanduser("~/Desktop")
            out_file = os.path.join(desktop_dir, f"{safe_title}_Summary.md")

            with open(out_file, "w", encoding="utf-8") as sf:
                sf.write(f"# YouTube Summary: {video_title}\n\n")
                sf.write(f"**Channel:** {uploader} | **Link:** [{video_url}]({video_url})\n\n")
                sf.write(f"*Generated by Swan AI on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n\n")
                sf.write("---\n\n")
                sf.write(summary_content)
                sf.write("\n")

            try:
                subprocess.Popen(["open", out_file])
            except Exception:
                pass

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = f"YouTube video summarized: {out_file}"
            self._sync_hud(recent_title=f"{safe_title[:20]} summarized", is_completion=True)
            self._play_chime()

        except Exception as e:
            print(f"❌ [YouTube Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="YouTube summary failed", is_completion=True)

    def _run_document_worker(self, task: AgentTask, title: str, content: str, format: str = "docx", file_name: str = ""):
        fmt = format.lower().strip()
        if fmt not in ("docx", "pdf", "markdown", "md", "txt"):
            fmt = "docx"
        if fmt == "md":
            fmt = "markdown"

        print(f"📄 [Document Agent] Generating {fmt.upper()} document: '{title}'", flush=True)
        _update_hud_working(f"Creating {fmt.upper()} 📄", f"{title[:25]}...")
        try:
            desktop_dir = os.path.expanduser("~/Desktop")
            safe_name = file_name.strip() if file_name else "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
            if not safe_name:
                safe_name = f"Swan_Document_{int(time.time())}"

            if len(content) < 500:
                try:
                    from google import genai
                    client = genai.Client(api_key=config.api_key)
                    enrich_prompt = (
                        f"You are a professional executive document author.\n"
                        f"Title: {title}\n"
                        f"Topic/Instructions: {content}\n\n"
                        "Expand this into a well-structured, comprehensive professional document with clear headings (#, ##), structured bullet points, and high-impact executive prose."
                    )
                    res = self._generate_with_fallback(
                        client,
                        contents=enrich_prompt,
                        config=None,
                        models=("gemini-flash-latest", "gemini-3.1-flash-lite-preview", "gemini-flash-lite-latest", "gemini-2.5-flash")
                    )
                    if res and res.text:
                        content = res.text.strip()
                except Exception:
                    pass

            if fmt == "docx":
                out_path = os.path.join(desktop_dir, f"{safe_name}.docx")
                import docx
                from docx.shared import Inches, Pt, RGBColor
                from docx.enum.text import WD_ALIGN_PARAGRAPH

                doc = docx.Document()
                title_p = doc.add_paragraph()
                title_run = title_p.add_run(title)
                title_run.font.size = Pt(24)
                title_run.font.bold = True
                title_run.font.color.rgb = RGBColor(30, 41, 59)
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                meta_p = doc.add_paragraph()
                meta_run = meta_p.add_run(f"Generated by Swan AI Agent • {datetime.now().strftime('%B %d, %Y')}")
                meta_run.font.size = Pt(10)
                meta_run.font.italic = True
                meta_run.font.color.rgb = RGBColor(100, 116, 139)
                meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph()

                for line in content.split("\n"):
                    s = line.strip()
                    if not s:
                        continue
                    if s.startswith("# "):
                        doc.add_heading(s[2:], level=1)
                    elif s.startswith("## "):
                        doc.add_heading(s[3:], level=2)
                    elif s.startswith("### "):
                        doc.add_heading(s[4:], level=3)
                    elif s.startswith("- ") or s.startswith("* "):
                        doc.add_paragraph(s[2:], style="List Bullet")
                    elif s.startswith("1. ") or s.startswith("2. ") or s.startswith("3. "):
                        doc.add_paragraph(s[3:], style="List Number")
                    else:
                        doc.add_paragraph(s)

                doc.save(out_path)

            elif fmt == "pdf":
                out_path = os.path.join(desktop_dir, f"{safe_name}.pdf")
                from reportlab.lib.pagesizes import letter
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.lib import colors

                doc = SimpleDocTemplate(out_path, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
                styles = getSampleStyleSheet()

                title_style = ParagraphStyle(
                    'DocTitle',
                    parent=styles['Heading1'],
                    fontSize=22,
                    leading=26,
                    textColor=colors.HexColor('#0f172a'),
                    spaceAfter=8,
                    alignment=1
                )
                meta_style = ParagraphStyle(
                    'DocMeta',
                    parent=styles['Normal'],
                    fontSize=9,
                    leading=12,
                    textColor=colors.HexColor('#64748b'),
                    spaceAfter=18,
                    alignment=1
                )
                h2_style = ParagraphStyle(
                    'DocH2',
                    parent=styles['Heading2'],
                    fontSize=14,
                    leading=18,
                    textColor=colors.HexColor('#1e293b'),
                    spaceBefore=14,
                    spaceAfter=6
                )
                body_style = ParagraphStyle(
                    'DocBody',
                    parent=styles['BodyText'],
                    fontSize=10,
                    leading=15,
                    textColor=colors.HexColor('#334155'),
                    spaceAfter=8
                )

                story = [
                    Paragraph(title, title_style),
                    Paragraph(f"Generated by Swan AI • {datetime.now().strftime('%B %d, %Y')}", meta_style),
                    Spacer(1, 10)
                ]

                for line in content.split("\n"):
                    s = line.strip()
                    if not s:
                        continue
                    clean_s = s.replace("<", "&lt;").replace(">", "&gt;")
                    if s.startswith("# ") or s.startswith("## "):
                        h_text = clean_s.lstrip("# ").strip()
                        story.append(Paragraph(h_text, h2_style))
                    elif s.startswith("- ") or s.startswith("* "):
                        bullet_text = clean_s[2:].strip()
                        story.append(Paragraph(f"• {bullet_text}", body_style))
                    else:
                        story.append(Paragraph(clean_s, body_style))

                doc.build(story)

            else:
                out_path = os.path.join(desktop_dir, f"{safe_name}.md")
                with open(out_path, "w", encoding="utf-8") as mf:
                    mf.write(f"# {title}\n\n")
                    mf.write(f"*Generated by Swan AI • {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n\n")
                    mf.write("---\n\n")
                    mf.write(content)
                    mf.write("\n")

            try:
                subprocess.Popen(["open", out_path])
            except Exception:
                pass

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = f"{fmt.upper()} document created: {out_path}"
            self._sync_hud(recent_title=f"{safe_name[:20]} created", is_completion=True)
            self._play_chime()

        except Exception as e:
            print(f"❌ [Document Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="Document creation failed", is_completion=True)

    def _run_presentation_worker(self, task: AgentTask, title: str, topic_or_content: str, slide_count: int = 5):
        print(f"📊 [Presentation Agent] Generating {slide_count} slides: '{title}'", flush=True)
        _update_hud_working("Presentation Agent 📊", f"Designing {title[:22]}...")
        try:
            from google import genai
            from google.genai import types
            client = genai.Client(api_key=config.api_key)

            slide_prompt = (
                f"You are a master presentation designer and keynote architect.\n"
                f"Title: {title}\n"
                f"Topic & Raw Material: {topic_or_content}\n"
                f"Slide Count: {slide_count}\n\n"
                "Design a high-impact, modern 16:9 presentation deck.\n"
                "Return a JSON array of slide objects. Each slide object MUST have:\n"
                "- 'title': str (short, memorable slide heading)\n"
                "- 'subtitle': str (optional category or tagline)\n"
                "- 'points': list of 3-4 bullet strings (high value, concise)\n"
                "- 'takeaway': str (one punchy highlight sentence)\n\n"
                "Output ONLY valid raw JSON array."
            )

            res = self._generate_with_fallback(
                client,
                contents=slide_prompt,
                config=types.GenerateContentConfig(response_mime_type="application/json"),
                models=("gemini-flash-latest", "gemini-3.1-flash-lite-preview", "gemini-flash-lite-latest", "gemini-2.5-flash")
            )

            slides_data = json.loads(res.text) if res and res.text else []
            if not isinstance(slides_data, list) or not slides_data:
                slides_data = [
                    {"title": title, "subtitle": "Executive Briefing", "points": [topic_or_content], "takeaway": "Key strategy overview"}
                ]

            desktop_dir = os.path.expanduser("~/Desktop")
            safe_name = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
            if not safe_name:
                safe_name = f"Swan_Presentation_{int(time.time())}"

            # 1. Generate PPTX
            pptx_path = os.path.join(desktop_dir, f"{safe_name}.pptx")
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor

            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank_layout = prs.slide_layouts[6]

            # Slide 1: Title Slide
            s1 = prs.slides.add_slide(blank_layout)
            bg = s1.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(15, 23, 42)
            bg.line.fill.background()

            t_box = s1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
            tf = t_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = RGBColor(248, 250, 252)

            sub_p = tf.add_paragraph()
            sub_p.text = slides_data[0].get("subtitle") or "Executive Strategy & Analysis"
            sub_p.font.size = Pt(20)
            sub_p.font.color.rgb = RGBColor(148, 163, 184)
            sub_p.space_before = Pt(14)

            # Content Slides
            for i, s_info in enumerate(slides_data):
                if i == 0 and len(slides_data) > 1:
                    continue
                slide = prs.slides.add_slide(blank_layout)
                c_bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
                c_bg.fill.solid()
                c_bg.fill.fore_color.rgb = RGBColor(248, 250, 252)
                c_bg.line.fill.background()

                h_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
                h_tf = h_box.text_frame
                h_p = h_tf.paragraphs[0]
                h_p.text = s_info.get("title", f"Section {i+1}")
                h_p.font.size = Pt(32)
                h_p.font.bold = True
                h_p.font.color.rgb = RGBColor(15, 23, 42)

                b_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.6))
                b_tf = b_box.text_frame
                b_tf.word_wrap = True

                for pt in s_info.get("points", []):
                    bp = b_tf.add_paragraph()
                    bp.text = f"• {pt}"
                    bp.font.size = Pt(20)
                    bp.font.color.rgb = RGBColor(51, 65, 85)
                    bp.space_after = Pt(14)

                takeaway = s_info.get("takeaway")
                if takeaway:
                    card = slide.shapes.add_shape(1, Inches(1.0), Inches(6.0), Inches(11.333), Inches(0.9))
                    card.fill.solid()
                    card.fill.fore_color.rgb = RGBColor(226, 232, 240)
                    card.line.fill.background()
                    c_tf = card.text_frame
                    cp = c_tf.paragraphs[0]
                    cp.text = f"💡 Key Takeaway: {takeaway}"
                    cp.font.size = Pt(14)
                    cp.font.bold = True
                    cp.font.color.rgb = RGBColor(30, 41, 59)

            prs.save(pptx_path)

            # 2. Companion HTML slides
            html_path = os.path.join(desktop_dir, f"{safe_name}_slides.html")
            html_slides = []
            for idx, s in enumerate(slides_data):
                pts_html = "".join(f"<li>{p}</li>" for p in s.get("points", []))
                takeaway_html = f"<div class='takeaway'><strong>Key Takeaway:</strong> {s.get('takeaway')}</div>" if s.get("takeaway") else ""
                html_slides.append(f"""
                <section class="slide" id="slide-{idx+1}">
                    <div class="slide-num">{idx+1} / {len(slides_data)}</div>
                    <h2>{s.get('title')}</h2>
                    <div class="subtitle">{s.get('subtitle', '')}</div>
                    <ul>{pts_html}</ul>
                    {takeaway_html}
                </section>
                """)

            full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{ margin: 0; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: #0f172a; color: #f8fafc; display: flex; flex-direction: column; align-items: center; padding: 40px 20px; }}
  .deck {{ width: 100%; max-width: 960px; display: flex; flex-direction: column; gap: 40px; }}
  .slide {{ background: #1e293b; border-radius: 16px; padding: 48px; box-shadow: 0 20px 40px rgba(0,0,0,0.4); border: 1px solid rgba(255,255,255,0.08); position: relative; min-height: 400px; display: flex; flex-direction: column; justify-content: center; }}
  .slide-num {{ position: absolute; top: 24px; right: 28px; font-size: 13px; color: #64748b; font-weight: 600; }}
  h2 {{ margin: 0 0 8px 0; font-size: 32px; color: #38bdf8; }}
  .subtitle {{ color: #94a3b8; font-size: 18px; margin-bottom: 24px; }}
  ul {{ margin: 0 0 24px 0; padding-left: 24px; font-size: 18px; line-height: 1.8; color: #e2e8f0; }}
  .takeaway {{ background: rgba(56, 189, 248, 0.1); border-left: 4px solid #38bdf8; padding: 14px 18px; border-radius: 6px; font-size: 15px; color: #bae6fd; }}
</style>
</head>
<body>
  <div class="deck">
    {''.join(html_slides)}
  </div>
</body>
</html>"""
            with open(html_path, "w", encoding="utf-8") as hf:
                hf.write(full_html)

            try:
                subprocess.Popen(["open", pptx_path])
            except Exception:
                pass

            task.status = "completed"
            task.finished_at = time.time()
            task.result_message = f"Presentation slides created: {pptx_path}"
            self._sync_hud(recent_title=f"{safe_name[:20]} slides ready", is_completion=True)
            self._play_chime()

        except Exception as e:
            print(f"❌ [Presentation Agent Error]: {e}", flush=True)
            task.status = "failed"
            task.finished_at = time.time()
            task.error = str(e)
            self._sync_hud(recent_title="Presentation failed", is_completion=True)

    def get_status(self, task_id: Optional[str] = None) -> Dict[str, Any]:
        with self._lock:
            if task_id and task_id in self.tasks:
                t = self.tasks[task_id]
                return {
                    "task_id": t.task_id,
                    "agent_type": t.agent_type,
                    "title": t.title,
                    "status": t.status,
                    "duration": f"{time.time() - t.started_at:.1f}s" if t.status == "running" else f"{t.finished_at - t.started_at:.1f}s",
                    "result": t.result_message,
                    "error": t.error
                }
            # Return summary of all recent tasks
            running_tasks = [t for t in self.tasks.values() if t.status == "running"]
            completed_tasks = [t for t in self.tasks.values() if t.status == "completed"]
            return {
                "active_agents_count": len(running_tasks),
                "running": [{"task_id": t.task_id, "title": t.title, "type": t.agent_type} for t in running_tasks],
                "recent_completed": [{"task_id": t.task_id, "title": t.title} for t in completed_tasks[-3:]]
            }

agent_manager = AgentManager()
